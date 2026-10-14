"""Bangkitkan src/lib/closing/pendapatanMitraClosingSeptember2026.ts dari sheet FIXED SEPTEMBER 2026.

Baris 53 (NET PROFIT SS) = bagian laba/fee manajemen untuk pusat, baris 54 (NET PROFIT HPP) =
margin HPP pusat; jumlah keduanya = baris 55 (TOTAL PROFIT SS) yang ikut di B75.
Pakai: python -I bangkitkan_pendapatan_mitra_sept2026.py <xlsx> <out.ts>
"""
import sys
import openpyxl

KOLOM = [  # kolom Excel, header baris 2, id outlet
    ('N', 'SS CIBINONG', '550e8400-e29b-41d4-a716-446655440014', 'MITRA CIBINONG'),
    ('O', 'SS CISEENG', '550e8400-e29b-41d4-a716-446655440017', 'MITRA CISEENG'),
    ('P', 'SS PEKAYON', '550e8400-e29b-41d4-a716-446655440018', 'MITRA PEKAYON'),
    ('Q', 'SS KALISARI', '550e8400-e29b-41d4-a716-446655440019', 'MITRA KALISARI'),
    ('R', 'SS SAWANGAN', '5a4df577-5237-476e-b54c-9eb642a5a516', 'MITRA SAWANGAN DTC'),
    ('S', 'SS SENTUL', '43b7bbd1-1fd4-44b5-87ca-b07a271151af', 'MITRA SENTUL'),
    ('T', 'SS KOTWIS', '3f38c41d-11e3-49ce-a189-d7303e45f9ad', 'MITRA CIBUBUR'),
    ('U', 'SS CICURUG', 'd9a2ef93-c298-4501-a471-1c5e2b3dff08', 'MITRA CICURUG'),
    ('V', 'SS PALEDANG', '550e8400-e29b-41d4-a716-446655440003', 'MITRA PALEDANG'),
    ('W', 'SS PAMULANG', 'bba67dba-2dca-4e98-bdb2-6a9e265e288c', 'MITRA PAMULANG'),
    ('X', 'SS CILEUNGSI', '62a56103-2085-4dd5-9d25-a3c0cffc88ff', 'MITRA CILEUNGSI'),
]

def main(xlsx, out):
    ws = openpyxl.load_workbook(xlsx, data_only=True)['FIXED SEPTEMBER 2026']
    for r, label in ((53, 'NET PROFIT SS'), (54, 'NET PROFIT HPP'), (55, 'TOTAL PROFIT SS')):
        if str(ws[f'A{r}'].value).strip() != label:
            sys.exit(f'A{r} bukan {label}: {ws[f"A{r}"].value!r}')
    baris = []
    total = 0.0
    for col, header, oid, nama in KOLOM:
        if str(ws[f'{col}2'].value).strip() != header:
            sys.exit(f'{col}2 bukan {header}: {ws[f"{col}2"].value!r}')
        ss = float(ws[f'{col}53'].value or 0)
        hpp = float(ws[f'{col}54'].value or 0)
        if abs(ss + hpp - float(ws[f'{col}55'].value or 0)) > 0.01:
            sys.exit(f'{col}53+{col}54 != {col}55')
        total += ss + hpp
        baris.append(f"    '{oid}': {{ nama: '{nama}', kolomExcel: '{col}', bagiHasilSs: {ss!r}, marginHpp: {hpp!r} }},")
    ts = f"""// DIBANGKITKAN oleh scripts/closing/bangkitkan_pendapatan_mitra_sept2026.py — jangan diedit tangan.
// Sumber: SS LAP. LABA RUGI 2026 fixxxxx.xlsx, sheet FIXED SEPTEMBER 2026, baris 53–54 kolom N–X.

/* Pendapatan pusat dari mitra versi closing September 2026. Excel menghitung
 * "TOTAL NET PROFIT INTERNAL" (B75) = laba 12 outlet internal + N55:X55, dengan
 * N55 = NET PROFIT SS (bagi hasil/fee manajemen) + NET PROFIT HPP (margin HPP).
 * Rumus ini berbeda dari hitungan live (fee 3% + 10% HPP dasar) dan memang
 * dipakai apa adanya untuk September saja. */

export interface PendapatanMitraClosing {{
  nama: string
  kolomExcel: string
  /** Baris 53 NET PROFIT SS. */
  bagiHasilSs: number
  /** Baris 54 NET PROFIT HPP. */
  marginHpp: number
}}

export const PENDAPATAN_MITRA_CLOSING_SEPT_2026 = {{
  from: '2026-09-01',
  to: '2026-09-30',
  sumber: 'SS LAP. LABA RUGI 2026 fixxxxx.xlsx — FIXED SEPTEMBER 2026 N53:X54',
  outlets: {{
{chr(10).join(baris)}
  }} as Record<string, PendapatanMitraClosing>,
}}

/** Total pendapatan pusat dari mitra versi closing, hanya untuk rentang tepat 1–30 September 2026. */
export function pendapatanMitraClosingUntuk(from: string, to: string): {{ feeManajemen: number; marginHpp: number }} | null {{
  const P = PENDAPATAN_MITRA_CLOSING_SEPT_2026
  if (from !== P.from || to !== P.to) return null
  let feeManajemen = 0
  let marginHpp = 0
  for (const o of Object.values(P.outlets)) {{
    feeManajemen += o.bagiHasilSs
    marginHpp += o.marginHpp
  }}
  return {{ feeManajemen, marginHpp }}
}}
"""
    open(out, 'w', encoding='utf-8', newline='\n').write(ts)
    print(f'OK: {len(KOLOM)} mitra, total {total:,.2f}')

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
