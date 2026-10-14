"""Bangkitkan src/lib/closing/labaKotorInternalClosingSeptember2026.ts dari sheet FIXED SEPTEMBER 2026.

Target per kolom internal B–M, dalam konvensi layar Laba Rugi (omzet kotor, TikTok kotor):
  omzet             = baris 3 + 9 + 16 + 24
  potongan          = baris 10 (admin food apps) + 18 (komisi TikTok GO)
  komisiTiktokKedua = baris 18 lagi — Excel memakai pencairan (baris 19) di baris 28 DAN
                      menambah baris 18 di baris 31, jadi komisi TikTok GO terpotong dua kali
  hpp = baris 29, waste = baris 30
Diperiksa: omzet − potongan − komisiTiktokKedua − hpp − waste = baris 33 tiap kolom.
Pakai: python -I bangkitkan_laba_kotor_internal_sept2026.py <xlsx> <out.ts>
"""
import sys
import openpyxl

KOLOM = [
    ('B', 'SS ONLINE', 'ss-online'), ('C', 'SS KITCHEN', '550e8400-e29b-41d4-a716-446655440001'),
    ('D', 'SS EMPANG', '550e8400-e29b-41d4-a716-446655440002'), ('E', 'SS BCC', '550e8400-e29b-41d4-a716-446655440004'),
    ('F', 'SS SUKMAJAYA', '550e8400-e29b-41d4-a716-446655440005'), ('G', 'SS JAGAKARSA', '550e8400-e29b-41d4-a716-446655440006'),
    ('H', 'SS BEJI', '550e8400-e29b-41d4-a716-446655440007'), ('I', 'SS SAWANGAN', '550e8400-e29b-41d4-a716-446655440008'),
    ('J', 'SS PAJAJARAN', '550e8400-e29b-41d4-a716-446655440009'), ('K', 'SS JATIWARINGIN', '550e8400-e29b-41d4-a716-446655440010'),
    ('L', 'SS CIRENDEU', '550e8400-e29b-41d4-a716-446655440011'), ('M', 'SS DRAMAGA', '550e8400-e29b-41d4-a716-446655440013'),
]
LABEL = {3: 'REVENUE', 9: 'REVENUE', 10: 'ADMIN FEE', 16: 'REVENUE', 18: 'ADMIN FEE (COM & CREATORS)',
         24: 'REVENUE', 29: 'TOTAL COGS (HPP)', 30: 'TOTAL WASTE', 33: 'TOTAL GROSS PROFIT'}

def main(xlsx, out):
    ws = openpyxl.load_workbook(xlsx, data_only=True)['FIXED SEPTEMBER 2026']
    for r, lab in LABEL.items():
        if str(ws[f'A{r}'].value).strip() != lab:
            sys.exit(f'A{r} bukan {lab}: {ws[f"A{r}"].value!r}')
    v = lambda c, r: float(ws[f'{c}{r}'].value or 0)
    baris, tot = [], dict(omzet=0.0, potongan=0.0, komisiTiktokKedua=0.0, hpp=0.0, waste=0.0)
    for col, header, oid in KOLOM:
        if str(ws[f'{col}2'].value).strip() != header:
            sys.exit(f'{col}2 bukan {header}: {ws[f"{col}2"].value!r}')
        o = dict(omzet=v(col, 3) + v(col, 9) + v(col, 16) + v(col, 24), potongan=v(col, 10) + v(col, 18),
                 komisiTiktokKedua=v(col, 18), hpp=v(col, 29), waste=v(col, 30))
        gp = o['omzet'] - o['potongan'] - o['komisiTiktokKedua'] - o['hpp'] - o['waste']
        if abs(gp - v(col, 33)) > 0.01:
            sys.exit(f'kolom {col}: hitungan {gp} != baris 33 {v(col, 33)}')
        for k in tot: tot[k] += o[k]
        isi = ', '.join(f'{k}: {o[k]!r}' for k in o)
        baris.append(f"    '{oid}': {{ kolomExcel: '{col}', {isi} }},")
    tot_isi = ', '.join(f'{k}: {tot[k]!r}' for k in tot)
    ts = f"""// DIBANGKITKAN oleh scripts/closing/bangkitkan_laba_kotor_internal_sept2026.py — jangan diedit tangan.
// Sumber: SS LAP. LABA RUGI 2026 fixxxxx.xlsx, sheet FIXED SEPTEMBER 2026, kolom B–M.

/* Target laba kotor outlet internal versi closing September 2026, dalam konvensi
 * layar Laba Rugi (omzet kotor, TikTok GO kotor). Termasuk keanehan Excel yang
 * dipilih owner: komisi TikTok GO terpotong dua kali (komisiTiktokKedua). */

export interface KomponenLabaKotor {{
  omzet: number
  potongan: number
  komisiTiktokKedua: number
  hpp: number
  waste: number
}}

export const LABA_KOTOR_INTERNAL_CLOSING_SEPT_2026 = {{
  from: '2026-09-01',
  to: '2026-09-30',
  sumber: 'SS LAP. LABA RUGI 2026 fixxxxx.xlsx — FIXED SEPTEMBER 2026 B3:M33',
  outlets: {{
{chr(10).join(baris)}
  }} as Record<string, KomponenLabaKotor & {{ kolomExcel: string }}>,
  total: {{ {tot_isi} }} as KomponenLabaKotor,
}}

/** Penyesuaian (target − live) per komponen, hanya untuk rentang tepat 1–30 September 2026.
 *  komisiTiktokKedua tidak punya padanan live, jadi nilainya = target penuh. */
export function penyesuaianLabaKotorClosing(
  from: string,
  to: string,
  live: {{ omzet: number; potongan: number; hpp: number; waste: number }},
): KomponenLabaKotor | null {{
  const L = LABA_KOTOR_INTERNAL_CLOSING_SEPT_2026
  if (from !== L.from || to !== L.to) return null
  return {{
    omzet: L.total.omzet - live.omzet,
    potongan: L.total.potongan - live.potongan,
    komisiTiktokKedua: L.total.komisiTiktokKedua,
    hpp: L.total.hpp - live.hpp,
    waste: L.total.waste - live.waste,
  }}
}}
"""
    open(out, 'w', encoding='utf-8', newline='\n').write(ts)
    print('OK: 12 kolom, ' + ', '.join(f'{k} {tot[k]:,.0f}' for k in tot))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
