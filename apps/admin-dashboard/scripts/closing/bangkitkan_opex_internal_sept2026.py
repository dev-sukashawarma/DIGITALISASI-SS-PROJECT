"""Bangkitkan data closing OPEX internal September 2026 dari sheet FIXED SEPTEMBER 2026.

Pakai (dari apps/admin-dashboard):
  python -I scripts/closing/bangkitkan_opex_internal_sept2026.py "<path xlsx>" src/lib/closing/opexClosingInternalSeptember2026.ts

File Excel TIDAK di-commit (berisi rincian gaji). Skrip berhenti kalau label baris
di kolom A tidak sesuai — artinya tata letak sheet berubah dan target harus dicek ulang.
"""
import json
import os
import sys

import openpyxl

SHEET = 'FIXED SEPTEMBER 2026'
OUTLETS = [
    ('B', 'ss-online', 'SS Online'),
    ('C', '550e8400-e29b-41d4-a716-446655440001', 'SUKA SHAWARMA BNR'),
    ('D', '550e8400-e29b-41d4-a716-446655440002', 'SUKA SHAWARMA EMPANG'),
    ('E', '550e8400-e29b-41d4-a716-446655440004', 'SUKA SHAWARMA CIMANGGU'),
    ('F', '550e8400-e29b-41d4-a716-446655440005', 'SUKA SHAWARMA DEPOK SUKMAJAYA'),
    ('G', '550e8400-e29b-41d4-a716-446655440006', 'SUKA SHAWARMA JAGAKARSA'),
    ('H', '550e8400-e29b-41d4-a716-446655440007', 'SUKA SHAWARMA BEJI'),
    ('I', '550e8400-e29b-41d4-a716-446655440008', 'SUKA SHAWARMA SAWANGAN (INTERNAL)'),
    ('J', '550e8400-e29b-41d4-a716-446655440009', 'SUKA SHAWARMA PAJAJARAN'),
    ('K', '550e8400-e29b-41d4-a716-446655440010', 'SUKA SHAWARMA JATIWARINGIN'),
    ('L', '550e8400-e29b-41d4-a716-446655440011', 'SUKA SHAWARMA CIRENDEU'),
    ('M', '550e8400-e29b-41d4-a716-446655440013', 'SUKA SHAWARMA DRAMAGA'),
]
BARIS_OUTLET = [
    (37, 'PENGELUARAN OUTLET', 'pengeluaran_outlet'),
    (38, 'GAJI CREW OUTLET', 'gaji_crew_outlet'),
    (39, 'BONUS CREW', 'bonus_crew'),
    (40, 'BONUS AREA MANAGER', 'bonus_area_manager'),
    (41, 'BONUS REGIONAL MANAGER', 'bonus_regional_manager'),
    (42, 'LEMBUR', 'lembur'),
    (43, 'ADS', 'ads'),
    (44, 'ENDORSEMENT', 'endorsement'),
    (45, 'PROMO', 'promo'),
    (46, 'PDAM', 'pdam'),
    (47, 'PLN', 'pln'),
    (48, 'INTERNET', 'internet'),
    (49, 'BIAYA SEWA OUTLET', 'sewa_outlet'),
]
TOTAL_OUTLET = (50, 'TOTAL OPEX')
BARIS_PUSAT = [
    (60, 'PDAM', 'pdam'),
    (61, 'PLN', 'pln'),
    (62, 'INTERNET', 'internet'),
    (65, 'PENGELUARAN GLOBAL', 'pengeluaran_global'),
    (66, 'GAJI STAFF KANTOR', 'gaji_staff_kantor'),
    (67, 'GAJI & BONUS CREW BACKUP', 'crew_backup'),
    (68, 'TRANSPORT', 'transport'),
    (69, 'BAHAN BAKU', 'bahan_baku'),
    (70, 'ADS', 'ads'),
]
TOTAL_PUSAT = (73, 'TOTAL OPEX HQ')

ALASAN_UMUM = {
    'pengeluaran_outlet': 'Kas kecil dikelompokkan ulang ke Pengeluaran Outlet sesuai closing',
    'gaji_crew_outlet': 'Kru cadangan dibebankan ke pusat; alokasi gaji AM versi closing',
    'bonus_crew': 'Bonus mengikuti slip payroll',
    'bonus_area_manager': 'Bonus AM & RM mengikuti closing (RM dibagi rata 1/21)',
    'bonus_regional_manager': 'Bonus RM dibagi rata 1/21 sesuai closing',
    'lembur': 'Lembur kas kecil dihitung di Pengeluaran Outlet sesuai closing',
}
ALASAN_KHUSUS = {
    ('ss-online', 'gaji_crew_outlet'): 'Slip Gilang, Ismansyah & Abdul Hanan dibebankan ke SS Online',
    ('550e8400-e29b-41d4-a716-446655440001', 'gaji_crew_outlet'): 'Slip Gilang & Ismansyah dibebankan ke SS Online',
    ('550e8400-e29b-41d4-a716-446655440001', 'lembur'): 'Lemburan a.n Gilang tidak dihitung; lembur Aby pindah ke pusat',
    ('550e8400-e29b-41d4-a716-446655440001', 'internet'): 'Pulsa HP Kitchen tidak dihitung',
    ('550e8400-e29b-41d4-a716-446655440005', 'lembur'): 'Bongkaran Sawangan tidak dihitung; lembur kas kecil di Pengeluaran Outlet',
}
ALASAN_PUSAT = {
    'gaji_staff_kantor': 'Slip Abdul Hanan dibebankan ke SS Online',
    'crew_backup': 'Gaji & bonus kru SS BACKUP dibebankan ke pusat',
    'internet': 'Pulsa KOL & Pulsa Markom tidak dihitung',
}


def angka(v):
    v = float(v or 0)
    r = round(v)
    return r if abs(v - r) < 0.005 else round(v, 2)


def cek_label(ws, baris, label):
    teks = str(ws.cell(baris, 1).value or '').strip().upper()
    if teks != label:
        sys.exit(f'Label baris {baris} = {teks!r}, seharusnya {label!r}. Tata letak sheet berubah.')


def main(sumber, tujuan):
    ws = openpyxl.load_workbook(sumber, data_only=True)[SHEET]
    for b, label, _ in BARIS_OUTLET + BARIS_PUSAT:
        cek_label(ws, b, label)
    cek_label(ws, *TOTAL_OUTLET)
    cek_label(ws, *TOTAL_PUSAT)

    outlets = {}
    for kol, oid, nama in OUTLETS:
        target = {key: angka(ws[f'{kol}{b}'].value) for b, _, key in BARIS_OUTLET}
        total = angka(ws[f'{kol}{TOTAL_OUTLET[0]}'].value)
        if abs(sum(target.values()) - total) >= 1:
            sys.exit(f'Kolom {kol}: jumlah baris {sum(target.values())} != total {total}')
        alasan = {}
        for _, _, key in BARIS_OUTLET:
            teks = ALASAN_KHUSUS.get((oid, key)) or ALASAN_UMUM.get(key)
            if teks:
                alasan[key] = teks
        outlets[oid] = {'nama': nama, 'kolomExcel': kol, 'target': target, 'totalExcel': total, 'alasan': alasan}

    target_pusat = {key: angka(ws[f'B{b}'].value) for b, _, key in BARIS_PUSAT}
    total_pusat = angka(ws[f'B{TOTAL_PUSAT[0]}'].value)
    if abs(sum(target_pusat.values()) - total_pusat) >= 1:
        sys.exit(f'HQ: jumlah baris {sum(target_pusat.values())} != total {total_pusat}')

    data = {
        'from': '2026-09-01',
        'to': '2026-09-30',
        'sumber': f'{os.path.basename(sumber)} — sheet {SHEET}',
        'outlets': outlets,
        'pusat': {'target': target_pusat, 'totalExcel': total_pusat, 'alasan': ALASAN_PUSAT},
    }
    isi = (
        '// File ini DIBANGKITKAN oleh scripts/closing/bangkitkan_opex_internal_sept2026.py.\n'
        '// Jangan diedit tangan. Salinan identik ada di apps/finance/src/lib/closing/.\n'
        '\n'
        "export type BarisOpexOutlet = 'pengeluaran_outlet' | 'gaji_crew_outlet' | 'bonus_crew' | 'bonus_area_manager' | 'bonus_regional_manager' | 'lembur' | 'ads' | 'endorsement' | 'promo' | 'pdam' | 'pln' | 'internet' | 'sewa_outlet'\n"
        "export type BarisOpexPusat = 'pdam' | 'pln' | 'internet' | 'pengeluaran_global' | 'gaji_staff_kantor' | 'crew_backup' | 'transport' | 'bahan_baku' | 'ads'\n"
        '\n'
        'export interface ClosingOpexOutlet {\n'
        '  nama: string\n  kolomExcel: string\n  target: Record<BarisOpexOutlet, number>\n'
        '  totalExcel: number\n  alasan: Partial<Record<BarisOpexOutlet, string>>\n}\n'
        'export interface ClosingOpexPusat {\n'
        '  target: Record<BarisOpexPusat, number>\n  totalExcel: number\n'
        '  alasan: Partial<Record<BarisOpexPusat, string>>\n}\n'
        'export interface ClosingOpexBulan {\n'
        '  from: string\n  to: string\n  sumber: string\n'
        '  outlets: Record<string, ClosingOpexOutlet>\n  pusat: ClosingOpexPusat\n}\n'
        '\n'
        f'export const BARIS_OPEX_OUTLET: BarisOpexOutlet[] = {json.dumps([k for _, _, k in BARIS_OUTLET])}\n'
        f'export const BARIS_OPEX_PUSAT: BarisOpexPusat[] = {json.dumps([k for _, _, k in BARIS_PUSAT])}\n'
        '\n'
        f'export const CLOSING_OPEX_INTERNAL_SEPT_2026: ClosingOpexBulan = {json.dumps(data, ensure_ascii=False, indent=2)}\n'
    )
    with open(tujuan, 'w', encoding='utf-8', newline='\n') as f:
        f.write(isi)
    print(f'OK: {len(outlets)} outlet, HQ {total_pusat} -> {tujuan}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
